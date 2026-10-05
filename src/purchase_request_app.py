from __future__ import annotations

import csv
import base64
import io
import json
import logging
import os
import re
import sqlite3
import threading
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from urllib import request as urllib_request

from dotenv import load_dotenv
from flask import Flask, jsonify, render_template, request
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parents[1]
load_dotenv(BASE_DIR / ".env", override=True)
os.environ.setdefault("EZACCOUNT_BASE_URL", "https://ezaccount.wamc.co.kr")
os.environ.setdefault("MASTER_REFRESH_URL", "")
DB_PATH = BASE_DIR / "state" / "purchase_requests.db"
UPLOAD_DIR = BASE_DIR / "uploads"
MASTER_FILE = BASE_DIR / "config" / "company_masters.json"
MASTER_REFRESH_URL = os.getenv("MASTER_REFRESH_URL", "").strip()
logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
REJECTION_POLL_SECONDS = max(60, int(os.getenv("EZACCOUNT_REJECTION_POLL_SECONDS", "300")))
_rejection_monitor_lock = threading.Lock()
_rejection_monitor_started = False


def get_ezaccount_credentials() -> Tuple[str, str]:
    load_dotenv(BASE_DIR / ".env", override=True)
    username = (os.getenv("EZACCOUNT_USERNAME") or "").strip()
    password = (os.getenv("EZACCOUNT_PASSWORD") or "").strip()
    if username == "your_user_id" or not username:
        username = ""
    if password == "your_password" or not password:
        password = ""
    return username, password


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def get_db_connection() -> sqlite3.Connection:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def normalize_master_rows(rows: List[Dict[str, Any]], kind: str) -> List[Dict[str, str]]:
    normalized: List[Dict[str, str]] = []
    seen: set[Tuple[str, str]] = set()

    for row in rows:
        if not isinstance(row, dict):
            continue
        values = {str(k).strip().lower(): str(v or "").strip() for k, v in row.items()}

        if kind == "vendor":
            code = (
                values.get("vendor_code") or values.get("vendorcode") or values.get("supplier_code") or values.get("suppliercode")
                or values.get("code") or values.get("vendor") or values.get("supplier")
            )
            name = (
                values.get("vendor_name") or values.get("vendorname") or values.get("supplier_name") or values.get("suppliername")
                or values.get("name") or values.get("company") or values.get("vendor text")
            )
        else:
            code = (
                values.get("material_code") or values.get("materialcode") or values.get("code") or values.get("material")
                or values.get("item_code") or values.get("itemcode")
            )
            name = (
                values.get("material_name") or values.get("materialname") or values.get("name") or values.get("description")
                or values.get("item_name") or values.get("itemname") or values.get("text")
            )

        code = str(code).strip()
        name = str(name).strip()
        if not code or not name:
            continue
        key = (code, name)
        if key in seen:
            continue
        seen.add(key)
        normalized.append({"code": code, "name": name})

    return normalized


def read_master_file(file_obj) -> List[Dict[str, Any]]:
    data = file_obj.read()
    name = (getattr(file_obj, "filename", "") or "").lower()

    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        return list(reader)

    if name.endswith(".xlsx") or name.endswith(".xls"):
        xls_bytes = io.BytesIO(data)
        wb = load_workbook(xls_bytes, read_only=True, data_only=True)
        ws = wb.active
        rows = []
        headers: List[str] = []
        if ws is None:
            raise ValueError("Worksheet is None")
        for row in ws.iter_rows(values_only=True):
            if not headers:
                headers = [str(v).strip() if v is not None else "" for v in row]
                continue
            rows.append({headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))})
        wb.close()
        return rows

    raise ValueError("지원하지 않는 파일 형식입니다. CSV 또는 Excel 파일만 가능합니다.")


def dedupe_master_items(items: List[Dict[str, str]]) -> List[Dict[str, str]]:
    seen: set[tuple[str, str]] = set()
    clean: List[Dict[str, str]] = []
    for item in items:
        if not isinstance(item, dict):
            continue
        code = str(item.get("code") or item.get("vendor_code") or item.get("material_code") or "").strip()
        name = str(item.get("name") or item.get("vendor_name") or item.get("material_name") or "").strip()
        if not code or not name:
            continue
        key = (code, name)
        if key in seen:
            continue
        seen.add(key)
        clean.append({"code": code, "name": name})
    return clean


def normalize_loaded(data: Dict[str, Any]) -> Dict[str, List[Dict[str, str]]]:
    vendors = data.get("vendors", []) if isinstance(data, dict) else []
    materials = data.get("materials", []) if isinstance(data, dict) else []
    return {
        "vendors": dedupe_master_items([
            {"code": str(v.get("code", "") or v.get("vendor_code", "") or v.get("supplier_code", "")).strip(), "name": str(v.get("name", "") or v.get("vendor_name", "") or v.get("supplier_name", "")).strip()}
            for v in vendors
        ]),
        "materials": dedupe_master_items([
            {"code": str(m.get("code", "") or m.get("material_code", "") or m.get("item_code", "")).strip(), "name": str(m.get("name", "") or m.get("material_name", "") or m.get("item_name", "")).strip()}
            for m in materials
        ]),
    }


def normalize_choice_rows(rows: List[Dict[str, Any]], kind: str) -> List[Dict[str, str]]:
    choices: List[Dict[str, str]] = []
    seen: set[Tuple[str, str]] = set()
    for row in rows:
        code = str(row.get("code") or row.get("value") or "").strip()
        name = str(row.get("name") or row.get("text") or "").strip()
        if not code and name:
            parts = name.split(None, 1)
            if len(parts) == 2:
                code, name = parts
        if not code or not name:
            continue
        code_pattern = {"vendors": r"\d{7}", "gl_accounts": r"\d{10}", "cost_centers": r"\d{7}"}[kind]
        if not re.fullmatch(code_pattern, code):
            continue
        key = (code, name)
        if key not in seen:
            seen.add(key)
            choices.append({"code": code, "name": name})
    return choices


def fetch_live_ezaccount_master_data() -> Dict[str, List[Dict[str, str]]]:
    try:
        from playwright.sync_api import sync_playwright
    except Exception:
        return {}

    try:
        username, password = get_ezaccount_credentials()
        if not username or not password:
            logging.warning("EZAccount master crawl skipped: credentials are not configured")
            return {}
        base_url = (os.getenv("EZACCOUNT_BASE_URL", "https://ezaccount.wamc.co.kr") or "https://ezaccount.wamc.co.kr").rstrip("/")
        login_url = f"{base_url}/sap/bc/ui5_ui5/sap/zeasy_account/index.html?sap-client=100"

        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            page = browser.new_page(viewport={"width": 1600, "height": 1100})
            page.on("dialog", lambda dialog: dialog.dismiss())

            def dismiss_login_popup() -> None:
                for _ in range(6):
                    dismissed = False
                    for selector in [
                        "#__button3", "button:has-text('Cancel')", "button:has-text('취소')",
                        "button:has-text('No')", "button:has-text('아니오')",
                    ]:
                        try:
                            candidate = page.locator(selector).first
                            if candidate.count() and candidate.is_visible():
                                candidate.click(timeout=3000, force=True)
                                page.wait_for_timeout(700)
                                dismissed = True
                                break
                        except Exception:
                            continue
                    if not dismissed:
                        break
                try:
                    page.wait_for_function(
                        "() => { const el = document.getElementById('sap-ui-blocklayer-popup'); return !el || getComputedStyle(el).display === 'none' || getComputedStyle(el).visibility === 'hidden'; }",
                        timeout=10000,
                    )
                except Exception:
                    pass

            page.goto(login_url, wait_until="domcontentloaded", timeout=60000)
            page.locator("#__xmlview1--USRID-inner").fill(username)
            page.locator("#__xmlview1--PASWD-inner").fill(password)
            page.locator("#__button0").click(timeout=20000)
            page.wait_for_timeout(6000)
            dismiss_login_popup()
            if page.locator("#__xmlview1--USRID-inner").is_visible():
                login_state = page.locator("body").inner_text()[:800]
                raise RuntimeError(f"EZAccount login did not leave the login screen: {login_state}")

            nav_candidates = [
                "div[title='Purchase Request']", "li[title='Purchase Request']",
                "#__item1-__xmlview2--NaviList-3-focusable", "#__item1-__xmlview2--NaviList-3",
                "#__item1-__xmlview2--NaviList-3-a",
                "#__item0-__xmlview2--NaviList-3-__item1-__xmlview2--NaviList-3-0",
                "#__item0-__xmlview2--NaviList-3-__item1-__xmlview2--NaviList-3-1",
            ]
            navigated = False
            start_url = page.url
            for selector in nav_candidates:
                try:
                    loc = page.locator(selector).first
                    if not loc.count():
                        continue
                    dismiss_login_popup()
                    loc.click(timeout=15000, force=True)
                    page.wait_for_timeout(1500)
                    if selector == "div[title='Purchase Request']":
                        nested = page.locator("li[title='Purchase Request']").first
                        if nested.count():
                            nested.click(timeout=15000, force=True)
                    page.wait_for_timeout(4000)
                    if page.url != start_url or "Purchase Request" in page.locator("body").inner_text():
                        navigated = True
                        break
                except Exception:
                    continue
            if not navigated:
                page_state = page.evaluate("""() => ({
                    body: (document.body?.innerText || '').slice(0, 1200),
                    titled: Array.from(document.querySelectorAll('[title]')).map(el => el.getAttribute('title')).filter(Boolean).slice(0, 80)
                })""")
                raise RuntimeError(f"Purchase Request menu was not reached (title={page.title()}, url={page.url}, state={page_state})")

            created = False
            for selector in ["button:has-text('New')", "button:has-text('신규')", "button[id*='create']", "#__button4", '[title="New"]']:
                try:
                    loc = page.locator(selector).first
                    if loc.count() > 0:
                        loc.click(timeout=20000, force=True)
                        page.wait_for_timeout(5000)
                        created = True
                        break
                except Exception:
                    continue
            if not created:
                raise RuntimeError(f"New PR form was not opened (title={page.title()}, url={page.url})")

            for selector in [
                "#__xmlview5--ObjectPageLayout-anchBar-__section1-anchor",
                "#__xmlview5--ADD_FAVORITE_APPROVAL", "#__xmlview5--BTN_APPROVER_LIST",
            ]:
                try:
                    section = page.locator(selector).first
                    if section.count():
                        section.click(timeout=5000, force=True)
                        page.wait_for_timeout(800)
                        for path in ["text=NEW PR", "text=New PR", "li:has-text('NEW PR')", "button:has-text('NEW PR')"]:
                            option = page.locator(path).first
                            if option.count():
                                option.click(timeout=5000, force=True)
                                page.wait_for_timeout(500)
                                break
                        break
                except Exception:
                    continue

            add_row = page.locator("#__button6").first
            if add_row.count():
                add_row.click(timeout=15000, force=True)
                page.wait_for_timeout(2000)
            collected: Dict[str, List[Dict[str, str]]] = {
                "vendors": [], "gl_accounts": [], "cost_centers": []
            }
            field_targets = {
                "vendors": "#__xmlview5--VENDOR-inner",
                "gl_accounts": "#__xmlview5--PURCHASE_DOCUMENT_GL-inner",
                "cost_centers": "#__xmlview5--AP_DOCUMENT_COST_CENTER-inner",
            }
            for kind, selector in field_targets.items():
                input_field = page.locator(selector).first
                if input_field.count() == 0:
                    logging.warning("EZAccount master field not found: %s", selector)
                    continue
                value_help_opened = False
                for value_help_selector in [selector.replace("-inner", "-vhi"), selector.replace("-inner", "-valueHelpIcon")]:
                    try:
                        value_help = page.locator(value_help_selector).first
                        if value_help.count() and value_help.is_visible():
                            value_help.click(timeout=3000, force=True)
                            page.wait_for_timeout(1200)
                            value_help_opened = True
                            break
                    except Exception:
                        continue
                if not value_help_opened:
                    try:
                        input_field.click(timeout=3000, force=True)
                        page.wait_for_timeout(700)
                    except Exception:
                        pass
                paged_rows: List[Dict[str, str]] = []
                if kind == "vendors":
                    search_field = page.locator(".sapMTableSelectDialog input[type='search']").first
                    vendor_row_selector = ".sapMTableSelectDialog .sapMListTblRow"

                    def read_vendor_rows() -> List[Dict[str, str]]:
                        return page.locator(vendor_row_selector).evaluate_all("""rows => rows
                            .filter(row => row.getBoundingClientRect().height > 0)
                            .map(row => {
                                const cells = Array.from(row.querySelectorAll('.sapMListTblCell'))
                                    .map(cell => (cell.getAttribute('title') || cell.getAttribute('aria-label') || cell.textContent || '').trim())
                                    .filter(Boolean);
                                return cells.length > 1 ? {code: cells[0], name: cells[1]} : null;
                            }).filter(Boolean)""")

                    def collect_vendor_prefix(prefix: str) -> None:
                        previous_text = page.locator(".sapMTableSelectDialog").inner_text()
                        search_field.fill(prefix)
                        search_field.press("Enter")
                        try:
                            page.wait_for_function(
                                "previous => { const dialog = document.querySelector('.sapMTableSelectDialog'); return dialog && dialog.innerText !== previous; }",
                                arg=previous_text,
                                timeout=3000,
                            )
                        except Exception:
                            page.wait_for_timeout(250)
                        result_rows = read_vendor_rows()
                        paged_rows.extend(result_rows)
                        vendor_codes = {row["code"] for row in result_rows if re.fullmatch(r"\d{7}", row["code"])}
                        if len(vendor_codes) >= 20 and len(prefix) < 7:
                            for digit in "0123456789":
                                collect_vendor_prefix(prefix + digit)

                    initial_rows = read_vendor_rows()
                    paged_rows.extend(initial_rows)
                    for first_digit in "0123456789":
                        collect_vendor_prefix(first_digit)
                    try:
                        search_field.fill("")
                        search_field.press("Enter")
                    except Exception:
                        pass
                for _ in range(100):
                    more_button = page.locator(".sapMTableSelectDialog button:has-text('More'), .sapMTableSelectDialog [role='button']:has-text('More')").first
                    if not more_button.count() or not more_button.is_visible():
                        break
                    table_dialog = page.locator(".sapMTableSelectDialog").first
                    previous_text = table_dialog.inner_text()
                    paged_rows.extend(page.evaluate("""() => {
                        const dialog = document.querySelector('.sapMTableSelectDialog');
                        if (!dialog) return [];
                        return Array.from(dialog.querySelectorAll('.sapMListTblRow')).map(row => {
                            const cells = Array.from(row.querySelectorAll('.sapMListTblCell'))
                                .map(cell => (cell.getAttribute('title') || cell.getAttribute('aria-label') || cell.textContent || '').trim())
                                .filter(Boolean);
                            return cells.length > 1 ? {code: cells[0], name: cells[1]} : null;
                        }).filter(Boolean);
                    }"""))
                    more_button.click(timeout=3000, force=True)
                    try:
                        page.wait_for_function(
                            "previous => { const dialog = document.querySelector('.sapMTableSelectDialog'); return !dialog || dialog.innerText !== previous; }",
                            arg=previous_text,
                            timeout=5000,
                        )
                    except Exception:
                        break
                rows = page.evaluate("""
                    () => {
                        const visible = el => {
                            const style = getComputedStyle(el);
                            return style.display !== 'none' && style.visibility !== 'hidden' && el.getBoundingClientRect().width > 0;
                        };
                        const overlays = Array.from(document.querySelectorAll('.sapMDialog, .sapMPopover, .sapMSelectList, [role="listbox"]')).filter(visible);
                        const roots = overlays.length ? overlays : [document];
                        const results = [];
                        const seen = new Set();
                        const add = (code, name) => {
                            code = (code || '').trim(); name = (name || '').trim();
                            if (!code && !name) return;
                            const key = `${code}::${name}`;
                            if (!seen.has(key)) { seen.add(key); results.push({code, name}); }
                        };
                        for (const root of roots) {
                            for (const select of root.querySelectorAll('select')) {
                                if (!visible(select)) continue;
                                for (const option of select.options) add(option.value, option.textContent);
                            }
                        }
                        const selectors = [
                            '[role="option"]', '.sapMSelectListItem', '.sapMComboBoxBaseItem',
                            '.sapMListTblRow', '.sapMLIB.sapMListTblRow', '.sapMListItems > .sapMLIB'
                        ];
                        for (const selector of selectors) {
                            for (const root of roots) for (const item of root.querySelectorAll(selector)) {
                                if (!visible(item)) continue;
                                if ((item.matches('.sapMListTblRow, .sapMLIB.sapMListTblRow') || item.closest('.sapMListTblRow'))
                                    && !item.closest('.sapMDialog, .sapMPopover, [role="listbox"]')) continue;
                                const cells = Array.from(item.querySelectorAll('td, .sapMListTblCell'))
                                    .map(cell => (cell.getAttribute('title') || cell.getAttribute('aria-label') || cell.textContent || '').trim()).filter(Boolean);
                                if (cells.length > 1) add(cells[0], cells[1]);
                                else {
                                    const text = (item.innerText || item.textContent || '').trim();
                                    const match = text.match(/^(\\S+)\\s+(.+)$/);
                                    if (match) add(match[1], match[2]);
                                }
                            }
                        }
                        const fieldId = location.hash;
                        const fieldIds = [fieldId];
                        const core = window.sap?.ui?.getCore?.();
                        for (const id of fieldIds) {
                            const control = core?.byId(id);
                            if (!control) continue;
                            const items = [
                                ...(control.getItems?.() || []),
                                ...(control.getSuggestionItems?.() || [])
                            ];
                            for (const item of items) {
                                add(item.getKey?.() || item.getAdditionalText?.() || '', item.getText?.() || '');
                            }
                        }
                        return results.slice(0, 5000);
                    }
                """.replace("location.hash", json.dumps(selector.replace("-inner", "").replace("#__", "__"))))
                collected[kind] = normalize_choice_rows(paged_rows + rows, kind)
                logging.info("Collected EZAccount %s choices: %s", kind, len(collected[kind]))
                try:
                    close_button = page.locator(".sapMDialog button[aria-label='Close'], .sapMDialog .sapMDialogCloseButton").first
                    if close_button.count() and close_button.is_visible():
                        close_button.click(timeout=2000, force=True)
                    else:
                        page.keyboard.press("Escape")
                    page.wait_for_timeout(500)
                except Exception:
                    continue
            browser.close()
            if not any(collected.values()):
                raise RuntimeError("PR form opened but no master choices were found")
            return collected
    except Exception:
        logging.exception("Unable to crawl EZAccount master choices")
        return {}


def load_company_master_data() -> Dict[str, List[Dict[str, str]]]:
    return fetch_live_ezaccount_master_data()


def ensure_master_columns(conn: sqlite3.Connection) -> None:
    for table, col_name in [("vendors", "updated_at"), ("materials", "updated_at")]:
        columns = [row[1] for row in conn.execute(f"PRAGMA table_info({table})").fetchall()]
        if col_name not in columns:
            conn.execute(f"ALTER TABLE {table} ADD COLUMN {col_name} TEXT")


def sync_choice_source(conn: sqlite3.Connection, table: str, rows: List[Dict[str, str]]) -> int:
    columns = {"gl_accounts": ("gl_code", "gl_name"), "cost_centers": ("cost_center_code", "cost_center_name")}
    code_key, name_key = columns[table]
    updated = 0
    conn.execute(f"DELETE FROM {table}")
    for item in rows:
        code = str(item.get("code") or "").strip()
        name = str(item.get("name") or "").strip()
        if code and name:
            conn.execute(f"INSERT OR IGNORE INTO {table} ({code_key}, {name_key}) VALUES (?, ?)", (code, name))
            updated += 1
    conn.commit()
    return updated


def sync_master_source(conn: sqlite3.Connection, kind: str, rows: List[Dict[str, str]]) -> int:
    ensure_master_columns(conn)

    if kind == "vendor":
        table = "vendors"
        code_key = "vendor_code"
        name_key = "vendor_name"
    else:
        table = "materials"
        code_key = "material_code"
        name_key = "material_name"

    inserted = 0
    for item in rows:
        code = str(item.get("code", "")).strip()
        name = str(item.get("name", "")).strip()
        if not code or not name:
            continue

        existing = conn.execute(
            f"SELECT 1 FROM {table} WHERE {code_key} = ?",
            (code,),
        ).fetchone()
        if existing:
            conn.execute(
                f"UPDATE {table} SET {name_key} = ?, updated_at = ? WHERE {code_key} = ?",
                (name, utc_now(), code),
            )
        else:
            conn.execute(
                f"INSERT INTO {table}({code_key}, {name_key}, created_at, updated_at) VALUES (?, ?, ?, ?)",
                (code, name, utc_now(), utc_now()),
            )
            inserted += 1

    conn.commit()
    return inserted


def ensure_pr_item_columns(conn: sqlite3.Connection) -> None:
    columns = [row[1] for row in conn.execute("PRAGMA table_info(pr_items)").fetchall()]
    if "material_code" not in columns:
        conn.execute("ALTER TABLE pr_items ADD COLUMN material_code TEXT")
    if "material_name" not in columns:
        conn.execute("ALTER TABLE pr_items ADD COLUMN material_name TEXT")
    if "manufacturing_part_no" not in columns:
        conn.execute("ALTER TABLE pr_items ADD COLUMN manufacturing_part_no TEXT")


def init_db() -> None:
    conn = get_db_connection()
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS vendors (
            vendor_code TEXT PRIMARY KEY,
            vendor_name TEXT NOT NULL,
            created_at TEXT NOT NULL DEFAULT (strftime('%Y-%m-%dT%H:%M:%SZ', 'now')),
            updated_at TEXT
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS materials (
            material_code TEXT PRIMARY KEY,
            material_name TEXT NOT NULL,
            created_at TEXT NOT NULL DEFAULT (strftime('%Y-%m-%dT%H:%M:%SZ', 'now')),
            updated_at TEXT
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS gl_accounts (
            gl_code TEXT PRIMARY KEY,
            gl_name TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS cost_centers (
            cost_center_code TEXT PRIMARY KEY,
            cost_center_name TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS master_refresh_state (
            id INTEGER PRIMARY KEY CHECK (id = 1),
            refreshed_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS master_import_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            kind TEXT NOT NULL,
            source_name TEXT NOT NULL,
            inserted_count INTEGER NOT NULL,
            imported_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS purchase_requests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            pr_number TEXT,
            status TEXT NOT NULL DEFAULT 'draft',
            revision INTEGER NOT NULL DEFAULT 0,
            total_amount REAL NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            ez_doc_no TEXT,
            ez_created_at TEXT,
            ez_status TEXT,
            notes TEXT
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS pr_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            request_id INTEGER NOT NULL,
            quote_group TEXT NOT NULL,
            item_name TEXT NOT NULL,
            purchase_reason TEXT,
            quantity REAL,
            unit_price REAL,
            total_price REAL,
            item_link TEXT,
            vendor_code TEXT,
            vendor_name TEXT,
            material_code TEXT,
            material_name TEXT,
            material_group TEXT,
            manufacturing_part_no TEXT,
            gl_account TEXT,
            cost_center TEXT,
            file_ref TEXT,
            created_at TEXT NOT NULL,
            FOREIGN KEY(request_id) REFERENCES purchase_requests(id)
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS revisions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            request_id INTEGER NOT NULL,
            revision INTEGER NOT NULL,
            summary TEXT NOT NULL,
            changed_at TEXT NOT NULL,
            FOREIGN KEY(request_id) REFERENCES purchase_requests(id)
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS attachments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            request_id INTEGER NOT NULL,
            item_id INTEGER,
            quote_group TEXT,
            display_order INTEGER,
            original_name TEXT NOT NULL,
            display_name TEXT,
            saved_name TEXT NOT NULL,
            saved_path TEXT NOT NULL,
            content_type TEXT,
            uploaded_at TEXT NOT NULL,
            FOREIGN KEY(request_id) REFERENCES purchase_requests(id)
        )
        """
    )

    ensure_pr_item_columns(conn)
    ensure_master_columns(conn)
    ensure_attachment_columns(conn)
    conn.commit()
    conn.close()


def seed_sample_data() -> None:
    conn = get_db_connection()
    existing = conn.execute("SELECT COUNT(*) FROM purchase_requests").fetchone()[0]
    if existing > 0:
        conn.close()
        return

    now = utc_now()
    conn.execute(
        "INSERT INTO purchase_requests (pr_number, status, revision, total_amount, created_at, updated_at, notes) VALUES (?, 'draft', 0, 0, ?, ?, 'MVP prototype seed data')",
        ("PR-INIT-001", now, now),
    )
    request_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
    conn.execute(
        "INSERT INTO pr_items (request_id, quote_group, item_name, purchase_reason, quantity, unit_price, total_price, item_link, vendor_code, vendor_name, material_code, material_name, material_group, gl_account, cost_center, file_ref, created_at) VALUES (?, 'quote-a', 'Motor Repair', 'Emergency repair', 1, 3200, 3200, 'https://example.com', '1100499', 'My Global LLC', '0007310602', 'CG_Maintenance_Equipment', 'MATERIAL', 'G/L 6000', 'CC-010', 'sample.pdf', ?)",
        (request_id, now),
    )
    conn.execute(
        "INSERT INTO pr_items (request_id, quote_group, item_name, purchase_reason, quantity, unit_price, total_price, item_link, vendor_code, vendor_name, material_code, material_name, material_group, gl_account, cost_center, file_ref, created_at) VALUES (?, 'quote-a', 'Labor', 'Emergency repair', 1, 4100, 4100, 'https://example.com', '1100499', 'My Global LLC', '0007310603', 'CG_Safety_Equipment', 'SERVICE', 'G/L 6000', 'CC-010', 'sample.pdf', ?)",
        (request_id, now),
    )
    conn.commit()
    conn.close()


def suggest_batches(items: List[Dict[str, Any]], cap: float = 10000.0) -> List[Dict[str, Any]]:
    groups: Dict[str, Dict[str, Any]] = {}
    for item in items:
        group_key = str(item.get("quote_group") or item.get("quoteGroup") or "ungrouped").strip() or "ungrouped"
        item_tot = item_total(item)
        groups.setdefault(group_key, {"key": group_key, "items": [], "total": 0.0})
        groups[group_key]["items"].append(item)
        groups[group_key]["total"] += item_tot

    sorted_groups = sorted(groups.values(), key=lambda g: g["total"], reverse=True)
    batches: List[Dict[str, Any]] = []
    for group in sorted_groups:
        placed = False
        for batch in batches:
            if batch["total"] + group["total"] <= cap:
                batch["groups"].append(group)
                batch["total"] += group["total"]
                placed = True
                break
        if not placed:
            batches.append({"groups": [group], "total": group["total"]})

    return batches


def item_total(item: Dict[str, Any]) -> float:
    try:
        qty = float(item.get("quantity") or 0)
        unit_price = float(item.get("unit_price") or item.get("unitPrice") or 0)
        total = qty * unit_price
        if total == 0.0:
            return float(item.get("total_price") or item.get("totalPrice") or 0.0)
        return total
    except Exception:
        return float(item.get("total_price") or item.get("totalPrice") or 0.0)


def save_uploaded_files(request_id: int, files: List[Any], manifest: Optional[List[Dict[str, Any]]] = None) -> None:
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    conn = get_db_connection()
    manifest = manifest or []
    for index, uploaded in enumerate(files):
        if uploaded is None or uploaded.filename == "":
            continue
        metadata = manifest[index] if index < len(manifest) else {}
        order = int(metadata.get("display_order") or index + 1)
        quote_group = str(metadata.get("quote_group") or "").strip()
        original_name = re.sub(r"[\x00-\x1f<>:\"/\\|?*]", "_", str(uploaded.filename).replace("\\", "/").rsplit("/", 1)[-1]).strip()
        original_name = original_name or f"attachment_{order}"
        display_name = f"{order}_{original_name}"
        extension = Path(original_name).suffix
        saved_name = f"{uuid.uuid4().hex}{extension}"
        saved_path = UPLOAD_DIR / saved_name
        uploaded.save(saved_path)
        conn.execute(
            "INSERT INTO attachments (request_id, quote_group, display_order, original_name, display_name, saved_name, saved_path, content_type, uploaded_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (request_id, quote_group, order, original_name, display_name, saved_name, str(saved_path), uploaded.mimetype or "application/octet-stream", utc_now()),
        )
    conn.commit()
    conn.close()


def ensure_attachment_columns(conn: sqlite3.Connection) -> None:
    columns = {row[1] for row in conn.execute("PRAGMA table_info(attachments)").fetchall()}
    for name, sql_type in (("quote_group", "TEXT"), ("display_order", "INTEGER"), ("display_name", "TEXT")):
        if name not in columns:
            conn.execute(f"ALTER TABLE attachments ADD COLUMN {name} {sql_type}")


def create_revision(request_id: int, revision: int, summary: str) -> None:
    conn = get_db_connection()
    conn.execute(
        "INSERT INTO revisions (request_id, revision, summary, changed_at) VALUES (?, ?, ?, ?)",
        (request_id, revision, summary, utc_now()),
    )
    conn.commit()
    conn.close()


def build_ezaccount_draft(request_id: int) -> Dict[str, Any]:
    conn = get_db_connection()
    request_row = conn.execute(
        "SELECT * FROM purchase_requests WHERE id = ?", (request_id,)
    ).fetchone()
    items = conn.execute(
        "SELECT * FROM pr_items WHERE request_id = ? ORDER BY id", (request_id,)
    ).fetchall()
    attachments = conn.execute(
        "SELECT * FROM attachments WHERE request_id = ? ORDER BY display_order, id", (request_id,)
    ).fetchall()
    conn.close()

    if request_row is None:
        raise ValueError(f"PR {request_id} not found")

    lines: List[Dict[str, Any]] = []
    for row in items:
        lines.append(
            {
                "item_name": row["item_name"],
                "purchase_reason": row["purchase_reason"],
                "quantity": row["quantity"],
                "unit_price": row["unit_price"],
                "total_price": row["total_price"],
                "vendor_code": row["vendor_code"],
                "vendor_name": row["vendor_name"],
                "material_code": row["material_code"],
                "material_name": row["material_name"],
                "material_group": row["material_group"],
                "manufacturing_part_no": row["manufacturing_part_no"],
                "gl_account": row["gl_account"],
                "cost_center": row["cost_center"],
                "reference_link": row["item_link"],
                "quote_group": row["quote_group"],
            }
        )

    return {
        "request_id": request_id,
        "status": "draft_ready",
        "header": {
            "document_type": "Purchase Request",
            "submit_on_user_click": True,
            "login_required": True,
            "user_session_strategy": "persistent_chrome_profile",
        },
        "items": lines,
        "attachments": [
            {
                "quote_group": row["quote_group"],
                "original_name": row["original_name"],
                "display_name": row["display_name"] or row["original_name"],
                "saved_path": row["saved_path"],
                "content_type": row["content_type"],
                "display_order": row["display_order"] or index,
            }
            for index, row in enumerate(attachments, start=1)
        ],
        "vendor_popup_flow": [
            "Open Supplier popup",
            "Enter Vendor Code",
            "Search vendor",
            "Select matched vendor record",
        ],
        "field_mapping": {
            "item_name": "Short Text",
            "purchase_reason": "Item Description",
            "quantity": "Qty",
            "unit_price": "Unit Price",
            "vendor_name": "Supplier",
            "material_code": "Material Code",
            "material_name": "Material Name",
            "material_group": "Material Group",
            "gl_account": "G/L Account",
            "cost_center": "Cost Center",
            "item_link": "Reference Link",
        },
        "playwright_steps": [
            "open EZAccount login",
            "reuse persistent browser profile",
            "create Purchase Request",
            "fill header fields",
            "repeat detail lines",
            "upload attachments",
            "save draft",
            "capture document number",
        ],
    }


def run_ezaccount_automation(request_id: int) -> Dict[str, Any]:
    load_dotenv(BASE_DIR / ".env", override=True)
    username, password = get_ezaccount_credentials()
    base_url = (os.getenv("EZACCOUNT_BASE_URL", "https://ezaccount.wamc.co.kr") or "https://ezaccount.wamc.co.kr").rstrip("/")

    if not username or not password:
        logging.warning("EZAccount automation skipped: missing credentials")
        raise RuntimeError("EZACCOUNT_USERNAME / EZACCOUNT_PASSWORD 환경 변수가 설정되지 않았습니다.")

    try:
        from playwright.sync_api import sync_playwright
    except Exception as exc:  # pragma: no cover
        logging.exception("Playwright unavailable")
        raise RuntimeError(f"Playwright가 설치되지 않았습니다: {exc}") from exc

    login_url = f"{base_url}/sap/bc/ui5_ui5/sap/zeasy_account/index.html?sap-client=100"
    logging.info("Starting EZAccount automation for request_id=%s base_url=%s login_url=%s", request_id, base_url, login_url)
    draft = build_ezaccount_draft(request_id)
    first_item = (draft.get("items") or [{}])[0]
    page_title = first_item.get("item_name") or "Purchase Request"

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page()

        def dismiss_any_popup() -> None:
            cancel_candidates = [
                "#__button3",
                "button:has-text('Cancel')",
                "button:has-text('취소')",
                "button:has-text('No')",
                "button:has-text('아니오')",
                "button:has-text('Close')",
                "button:has-text('닫기')",
                "button:has-text('OK')",
                "button:has-text('확인')",
                "//*[contains(text(),'Cancel')]",
                "//*[contains(text(),'취소')]",
                "//*[contains(text(),'No')]",
                "//*[contains(text(),'아니오')]",
            ]

            for _ in range(8):
                found = False
                for selector in cancel_candidates:
                    try:
                        if selector.startswith("#"):
                            loc = page.locator(selector).first
                        elif selector.startswith("//"):
                            loc = page.locator(f"xpath={selector}").first
                        else:
                            loc = page.locator(selector).first

                        if loc.count() > 0:
                            logging.info("Found cancel candidate selector=%s", selector)
                            loc.click(timeout=10000, force=True)
                            logging.info("Dismissed popup via selector=%s", selector)
                            found = True
                            break
                    except Exception:
                        continue
                if not found:
                    break
                page.wait_for_timeout(1000)

            try:
                page.wait_for_function(
                    """
                    () => {
                        const block = document.getElementById('sap-ui-blocklayer-popup');
                        if (!block) return true;
                        const style = window.getComputedStyle(block);
                        return style.display === 'none' || style.visibility === 'hidden' || style.opacity === '0';
                    }
                    """,
                    timeout=20000,
                )
                logging.info("SAP block layer cleared before continuing")
            except Exception:
                logging.warning("SAP block layer did not clear within timeout; continuing")

        def handle_dialog(dialog: Any) -> None:
            dialog.dismiss()
            logging.info("Dismissed browser dialog: %s", dialog.message)

        page.on("dialog", handle_dialog)

        logging.info("Opening EZAccount login page")
        page.goto(login_url, wait_until="domcontentloaded", timeout=60000)
        page.locator("#__xmlview1--USRID-inner").wait_for(timeout=30000)
        page.locator("#__xmlview1--PASWD-inner").wait_for(timeout=30000)

        login_selectors = [
            "#__xmlview1--USRID-inner",
            "input[name='j_username']",
            "input[name='username']",
            "input[id*='user']",
            "input[type='text']",
        ]
        password_selectors = [
            "#__xmlview1--PASWD-inner",
            "input[name='j_password']",
            "input[name='password']",
            "input[type='password']",
        ]

        filled = False
        for selector in login_selectors:
            try:
                loc = page.locator(selector).first
                if loc.count() > 0:
                    logging.info("Found login field selector=%s", selector)
                    loc.fill(username)
                    filled = True
                    break
            except Exception:
                continue
        if not filled:
            raise RuntimeError("EZAccount 로그인 화면에서 사용자 ID 입력 필드를 찾지 못했습니다.")

        for selector in password_selectors:
            try:
                loc = page.locator(selector).first
                if loc.count() > 0:
                    logging.info("Found password field selector=%s", selector)
                    loc.fill(password)
                    break
            except Exception:
                continue

        submit_selectors = [
            "#__button0",
            "button[type='submit']",
            "input[type='submit']",
            "button:has-text('Login')",
            "button:has-text('로그인')",
        ]
        clicked = False
        for selector in submit_selectors:
            try:
                loc = page.locator(selector).first
                if loc.count() > 0:
                    logging.info("Found login submit selector=%s", selector)
                    loc.click(timeout=20000)
                    clicked = True
                    break
            except Exception:
                continue
        if not clicked:
            raise RuntimeError("EZAccount 로그인 제출 버튼을 찾지 못했습니다.")

        page.wait_for_timeout(6000)
        dismiss_any_popup()
        target_url = f"{base_url}/sap/bc/ui5_ui5/sap/zeasy_account/index.html?sap-client=100"
        logging.info("Opening EZAccount app target=%s", target_url)
        page.goto(target_url, wait_until="domcontentloaded", timeout=60000)
        page.wait_for_timeout(4000)
        dismiss_any_popup()

        nav_selectors = [
            "div[title='Purchase Request']",
            "div[title=\"Purchase Request\"]",
            "li[title='Purchase Request']",
            "li[title=\"Purchase Request\"]",
            "#__item1-__xmlview2--NaviList-3-focusable",
            "#__item1-__xmlview2--NaviList-3",
            "#__item1-__xmlview2--NaviList-3-a",
            "#__item0-__xmlview2--NaviList-3-__item1-__xmlview2--NaviList-3-0",
            "#__item0-__xmlview2--NaviList-3-__item1-__xmlview2--NaviList-3-1",
        ]
        current_url = page.url
        pr_nav_ok = False
        for selector in nav_selectors:
            try:
                loc = page.locator(selector).first
                if loc.count() == 0:
                    continue
                logging.info("Trying EZAccount PR nav selector=%s", selector)
                dismiss_any_popup()
                page.wait_for_timeout(1000)
                loc.scroll_into_view_if_needed()
                page.wait_for_function(
                    """
                    () => {
                        const block = document.getElementById('sap-ui-blocklayer-popup');
                        return !block || getComputedStyle(block).display === 'none' || getComputedStyle(block).visibility === 'hidden';
                    }
                    """,
                    timeout=20000,
                )

                # SAP UI5 navigation is hierarchical: group item first, then nested Purchase Request item.
                if selector in ("div[title='Purchase Request']", "div[title=\"Purchase Request\"]"):
                    loc.click(timeout=20000, force=True)
                    page.wait_for_timeout(2000)
                    nested = page.locator("li[title='Purchase Request']").first
                    if nested.count() > 0:
                        nested.click(timeout=20000, force=True)
                else:
                    loc.click(timeout=20000, force=True)

                page.wait_for_timeout(5000)
                if page.url != current_url or page.locator("text=Purchase Request").count() > 0 or "PurchaseRequest" in page.url or "PurchaseDocument" in page.url:
                    pr_nav_ok = True
                    logging.info("PR navigation appears successful via selector=%s final_url=%s", selector, page.url)
                    break
            except Exception as exc:
                logging.warning("PR nav selector failed: %s -> %s", selector, exc)

        if not pr_nav_ok:
            raise RuntimeError("EZAccount 로그인은 성공했지만 Purchase Request 네비게이션이 동작하지 않았습니다. 실제 SAP PR 메뉴 selector를 더 매핑해야 합니다.")

        new_button_selectors = [
            "button:has-text('New')",
            "button:has-text('신규')",
            "button[id*='create']",
            "button[id*='add']",
            "[title='New']",
            "[title='Create']",
            "#__button4",
        ]
        new_button_clicked = False
        for selector in new_button_selectors:
            try:
                loc = page.locator(selector).first
                if loc.count() == 0:
                    continue
                logging.info("Clicking Purchase Request New button selector=%s", selector)
                dismiss_any_popup()
                page.wait_for_timeout(1000)
                loc.click(timeout=20000, force=True)
                new_button_clicked = True
                page.wait_for_timeout(5000)
                if "/PurchaseDocument/Create" in page.url or "Purchase Request Document Creation" in page.locator("body").inner_text():
                    logging.info("PR Create form opened: %s", page.url)
                    break
            except Exception as exc:
                logging.warning("New button selector failed: %s -> %s", selector, exc)

        if not new_button_clicked:
            raise RuntimeError("EZAccount 로그인과 PR 메뉴 이동은 성공했지만, 새 PR 생성용 New 버튼을 찾지 못했습니다.")

        approval_selectors = [
            "#__xmlview5--ObjectPageLayout-anchBar-__section1-anchor",
            "#__xmlview5--ADD_FAVORITE_APPROVAL",
            "#__xmlview5--BTN_APPROVER_LIST",
            "button:has-text('Search')",
            "text=NEW PR",
            "text=New PR",
            "button:has-text('NEW PR')",
            "button:has-text('New PR')",
            "li:has-text('NEW PR')",
            "div:has-text('NEW PR')",
        ]
        for selector in approval_selectors:
            try:
                loc = page.locator(selector).first
                if loc.count() > 0:
                    logging.info("Trying approval path selector=%s", selector)
                    loc.click(timeout=20000, force=True)
                    page.wait_for_timeout(1500)
                    for candidate in ["text=NEW PR", "text=New PR", "li:has-text('NEW PR')", "div:has-text('NEW PR')", "button:has-text('NEW PR')", "button:has-text('New PR')"]:
                        try:
                            target = page.locator(candidate).first
                            if target.count() > 0:
                                target.click(timeout=20000, force=True)
                                logging.info("Selected approval path via %s", candidate)
                                break
                        except Exception:
                            continue
                    break
            except Exception:
                continue

        first_item = (draft.get("items") or [{}])[0]
        item_name = first_item.get("item_name") or "Purchase Request for Press and Utility Materials"
        header_description = f"{item_name}, "
        header_note = first_item.get("purchase_reason") or "Purchase Request for Press and Utility Materials"
        for selector in ["#__xmlview5--HEADER_TEXT-inner", "#__xmlview5--HEADER_NOTE-inner", "textarea[title*='Header']"]:
            try:
                loc = page.locator(selector).first
                if loc.count() > 0:
                    loc.fill(header_description)
                    logging.info("Filled PR header description via selector=%s", selector)
                    break
            except Exception:
                continue

        if page.locator("#__xmlview5--HEADER_NOTE-inner").count() > 0 and page.locator("#__xmlview5--HEADER_TEXT-inner").count() > 0:
            try:
                page.locator("#__xmlview5--HEADER_NOTE-inner").fill(header_note)
            except Exception:
                pass

        item_count = 0
        for item in draft.get("items", []):
            if not item.get("item_name"):
                continue
            item_count += 1

            add_row_btn = page.locator("#__button6").first
            if add_row_btn.count() > 0:
                add_row_btn.click(timeout=20000, force=True)
                page.wait_for_timeout(3000)

            item_description = item.get("purchase_reason") or item.get("item_name") or ""
            item_description_detail = item.get("item_name") or item.get("purchase_reason") or ""
            item_text = f"1) {item_description}\n2) {item_description_detail}"
            field_map = {
                "#__xmlview5--VENDOR-inner": item.get("vendor_code") or item.get("vendor_name") or "",
                "#__xmlview5--PURCHASE_SHORT_TEXT-inner": item.get("item_name") or "",
                "#__xmlview5--PURCHASE_UNIT-inner": "EA",
                "#__xmlview5--PURCHASE_QTY-inner": item.get("quantity") or 0,
                "#__xmlview5--PURCHASE_UNIT_PRICE-inner": item.get("unit_price") or 0,
                "#__xmlview5--PURCHASE_TOTAL_PRICE-inner": item.get("total_price") or 0,
                "#__xmlview5--PURCHASE_PART_NO-inner": item.get("manufacturing_part_no") or "",
                "#__xmlview5--PURCHASE_DOCUMENT_GL-inner": item.get("gl_account") or "",
                "#__xmlview5--AP_DOCUMENT_COST_CENTER-inner": item.get("cost_center") or "",
                "#__xmlview5--ITEM_TEXT-inner": item_text,
            }

            if page.locator("#__xmlview5--ITEM_TEXT-inner").count() > 0:
                page.locator("#__xmlview5--ITEM_TEXT-inner").fill(item_text)
            if page.locator("#__xmlview5--PURCHASE_SHORT_TEXT-inner").count() > 0:
                page.locator("#__xmlview5--PURCHASE_SHORT_TEXT-inner").fill(str(item.get("item_name") or ""))
            for selector, value in field_map.items():
                try:
                    loc = page.locator(selector).first
                    if loc.count() > 0:
                        loc.fill(str(value))
                        page.wait_for_timeout(400)
                except Exception:
                    continue

            material_group = str(item.get("material_group") or "").strip()
            if material_group:
                group_value_help = page.locator("#__xmlview5--PURCHASE_ITEM_GROUP-vhi").first
                if not group_value_help.count():
                    raise RuntimeError("EZAccount Material Group value-help control was not found.")
                group_value_help.click(timeout=10000, force=True)
                page.wait_for_timeout(700)
                group_search = page.locator(".sapMTableSelectDialog input[type='search']").first
                if not group_search.count():
                    raise RuntimeError("EZAccount Material Group search dialog did not open.")
                group_search.fill(material_group)
                group_search.press("Enter")
                page.wait_for_timeout(700)
                group_rows = page.locator(".sapMTableSelectDialog .sapMListTblRow")
                selected_row = None
                for row_index in range(group_rows.count()):
                    row = group_rows.nth(row_index)
                    if material_group.casefold() in (row.inner_text() or "").casefold():
                        selected_row = row
                        break
                if selected_row is None:
                    raise RuntimeError(f"EZAccount Material Group was not found: {material_group}")
                selected_row.click(timeout=10000, force=True)
                page.wait_for_timeout(500)
                selected_group = page.locator("#__xmlview5--PURCHASE_ITEM_GROUP-inner").input_value()
                if material_group.casefold() not in selected_group.casefold():
                    raise RuntimeError(f"EZAccount did not select Material Group {material_group}.")

            reference_link = item.get("reference_link") or ""
            if reference_link:
                link_input = page.locator("#__xmlview5--LINK-inner").first
                if link_input.count() > 0:
                    link_input.fill(reference_link)
                    page.wait_for_timeout(400)
                    add_link_btn = page.locator("#__button8").first
                    if add_link_btn.count() > 0:
                        add_link_btn.click(timeout=20000, force=True)
                        page.wait_for_timeout(2000)

            item_save = page.locator("#__xmlview5--ITEM_SAVE").first
            if item_save.count() > 0:
                item_save.click(timeout=20000, force=True)
                page.wait_for_timeout(5000)
                logging.info("Saved PR detail row for %s", item.get("item_name"))
            else:
                logging.warning("PR detail row save button not found for item=%s", item.get("item_name"))

        attachments = draft.get("attachments") or []
        if attachments:
            if not page.locator("#__xmlview5--INVOICE_ATTACHMENT-uploader").count():
                raise RuntimeError("EZAccount PR attachment uploader was not found.")
            upload_files = []
            for attachment in attachments:
                saved_path = Path(str(attachment.get("saved_path") or ""))
                if not saved_path.is_file():
                    raise RuntimeError(f"Saved attachment is missing: {saved_path.name}")
                upload_files.append({
                    "name": str(attachment.get("display_name") or attachment.get("original_name") or saved_path.name),
                    "mime_type": str(attachment.get("content_type") or "application/octet-stream"),
                    "data": base64.b64encode(saved_path.read_bytes()).decode("ascii"),
                })
            staged = page.evaluate("""
                files => {
                    const uploadSet = sap.ui.getCore().byId('__xmlview5--INVOICE_ATTACHMENT');
                    if (!uploadSet) throw new Error('EZAccount UploadSet control is unavailable');
                    for (const entry of files) {
                        const binary = atob(entry.data);
                        const bytes = new Uint8Array(binary.length);
                        for (let index = 0; index < binary.length; index += 1) bytes[index] = binary.charCodeAt(index);
                        const file = new File([bytes], entry.name, {type: entry.mime_type});
                        const item = new sap.m.upload.UploadSetItem({
                            fileName: entry.name,
                            mediaType: entry.mime_type,
                            uploadState: 'Complete'
                        });
                        item._setFileObject(file);
                        uploadSet.addItem(item);
                    }
                    sap.ui.getCore().applyChanges();
                    return uploadSet.getItems().map(item => item.getFileName());
                }
            """, upload_files)
            if staged[-len(upload_files):] != [file["name"] for file in upload_files]:
                raise RuntimeError("EZAccount did not stage every shared quote attachment in order.")
            logging.info("Staged %s shared quote attachments in order: %s", len(upload_files), [file["name"] for file in upload_files])

        main_save = page.locator("#__button10").first
        if not main_save.count() or not main_save.is_visible() or (main_save.inner_text() or "").strip() != "Save":
            raise RuntimeError("EZAccount 하단 Save 버튼을 확인하지 못해 안전하게 중단했습니다. Submit은 실행하지 않았습니다.")
        main_save.click(timeout=20000, force=True)

        save_confirmation = page.locator("[role='alertdialog']").filter(has_text="Document will be saved. Do you want to continue?")
        try:
            save_confirmation.wait_for(state="visible", timeout=15000)
            save_confirmation.locator("button:has-text('OK')").click(timeout=10000, force=True)
        except Exception as exc:
            raise RuntimeError(f"EZAccount save confirmation was not completed. Submit was not pressed: {exc}") from exc

        save_success = page.locator("[role='alertdialog']").filter(has_text=re.compile(r"Request\s+\d{10}\s+is saved", re.IGNORECASE))
        try:
            save_success.wait_for(state="visible", timeout=30000)
            success_text = save_success.inner_text()
            success_match = re.search(r"\b\d{10}\b", success_text)
            document_no = success_match.group(0) if success_match else None
            save_success.locator("button:has-text('OK')").click(timeout=10000, force=True)
        except Exception as exc:
            raise RuntimeError(f"EZAccount did not confirm that the PR draft was saved. Submit was not pressed: {exc}") from exc

        for candidate in [
            "text=/PR[- ]\\d+/i",
            "div:has-text('Document Number')",
            "span:has-text('PR')",
        ]:
            try:
                loc = page.locator(candidate).first
                if loc.count() > 0:
                    text = loc.text_content(timeout=5000)
                    if text:
                        match = re.search(r"\b\d{10}\b", text)
                        if match:
                            document_no = match.group(0)
                            break
            except Exception:
                continue

        if not document_no:
            document_numbers = re.findall(r"\b\d{10}\b", page.locator("body").inner_text())
            if document_numbers:
                document_no = document_numbers[-1]

        browser.close()
        logging.info("EZAccount automation finished for request_id=%s document_no=%s item_count=%s", request_id, document_no, item_count)
        return {
            "ok": True,
            "request_id": request_id,
            "document_no": document_no,
            "status": "saved_draft",
            "submitted": False,
            "item_count": item_count,
            "title": page_title,
        }


def fetch_ezaccount_rejected_pr_numbers() -> set[str]:
    from playwright.sync_api import sync_playwright

    username, password = get_ezaccount_credentials()
    base_url = (os.getenv("EZACCOUNT_BASE_URL", "https://ezaccount.wamc.co.kr") or "https://ezaccount.wamc.co.kr").rstrip("/")
    login_url = f"{base_url}/sap/bc/ui5_ui5/sap/zeasy_account/index.html?sap-client=100"
    rejected_numbers: set[str] = set()

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1600, "height": 1100})
        page.on("dialog", lambda dialog: dialog.dismiss())
        page.goto(login_url, wait_until="domcontentloaded", timeout=60000)
        page.locator("#__xmlview1--USRID-inner").fill(username)
        page.locator("#__xmlview1--PASWD-inner").fill(password)
        page.locator("#__button0").click(timeout=20000)
        page.wait_for_timeout(6000)

        for _ in range(6):
            dismissed = False
            for selector in ["#__button3", "button:has-text('Cancel')", "button:has-text('취소')"]:
                try:
                    popup_action = page.locator(selector).first
                    if popup_action.count() and popup_action.is_visible():
                        popup_action.click(timeout=3000, force=True)
                        page.wait_for_timeout(500)
                        dismissed = True
                        break
                except Exception:
                    continue
            if not dismissed:
                break

        if page.locator("#__xmlview1--USRID-inner").is_visible():
            raise RuntimeError("EZAccount login did not leave the login screen while checking rejected PRs.")

        page.get_by_text("Purchase Request", exact=True).first.click(force=True, timeout=15000)
        page.wait_for_timeout(1000)
        page.get_by_text("PR List (Team)", exact=True).click(force=True, timeout=15000)
        page.wait_for_timeout(4000)

        date_filter = page.locator("#__xmlview4--POSTDATE-inner").first
        if date_filter.count():
            date_filter.fill("")
        status_filter = page.locator("#__xmlview4--STATUS").first
        status_filter.click(force=True, timeout=10000)
        page.get_by_text("Rejected", exact=True).click(force=True, timeout=10000)
        page.get_by_text("Go", exact=True).click(force=True, timeout=10000)
        page.wait_for_timeout(3000)

        for row_text in page.locator(".sapMListTblRow").all_text_contents():
            match = re.search(r"\b\d{10}\b", row_text)
            if match:
                rejected_numbers.add(match.group(0))

        browser.close()
    return rejected_numbers


def sync_rejected_purchase_requests() -> int:
    conn = get_db_connection()
    submitted = conn.execute(
        "SELECT id, ez_doc_no FROM purchase_requests WHERE status IN ('saved', 'submitted') AND COALESCE(ez_doc_no, '') != ''"
    ).fetchall()
    if not submitted:
        conn.close()
        return 0

    rejected_numbers = fetch_ezaccount_rejected_pr_numbers()
    changed = 0
    for row in submitted:
        if str(row["ez_doc_no"]).strip() not in rejected_numbers:
            continue
        conn.execute(
            "UPDATE purchase_requests SET status = 'rejected', ez_status = 'Rejected', updated_at = ? WHERE id = ?",
            (utc_now(), row["id"]),
        )
        changed += 1
    conn.commit()
    conn.close()
    if changed:
        logging.info("Queued %s rejected purchase requests for resubmission", changed)
    return changed


def rejection_monitor_loop() -> None:
    while True:
        try:
            sync_rejected_purchase_requests()
        except Exception:
            logging.exception("EZAccount rejected PR monitor failed")
        threading.Event().wait(REJECTION_POLL_SECONDS)


def start_rejection_monitor() -> None:
    global _rejection_monitor_started
    with _rejection_monitor_lock:
        if _rejection_monitor_started:
            return
        _rejection_monitor_started = True
        threading.Thread(target=rejection_monitor_loop, name="ezaccount-rejection-monitor", daemon=True).start()


def create_app() -> Flask:
    app = Flask(__name__, template_folder=str(BASE_DIR / "templates"), static_folder=str(BASE_DIR / "static"))
    app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024
    app.config["UPLOAD_FOLDER"] = str(UPLOAD_DIR)
    init_db()
    seed_sample_data()

    @app.route("/")
    def home() -> str:
        return render_template("purchase_requests.html")

    @app.route("/requests")
    def requests_page() -> str:
        return render_template("purchase_requests.html")

    @app.route("/api/vendors")
    def vendors_api():
        conn = get_db_connection()
        rows = conn.execute("SELECT vendor_code, vendor_name FROM vendors ORDER BY vendor_name").fetchall()
        conn.close()
        return jsonify({"items": [{"code": r["vendor_code"], "name": r["vendor_name"]} for r in rows]})

    @app.route("/api/materials")
    def materials_api():
        conn = get_db_connection()
        rows = conn.execute("SELECT material_code, material_name FROM materials ORDER BY material_name").fetchall()
        conn.close()
        return jsonify({"items": [{"code": r["material_code"], "name": r["material_name"]} for r in rows]})

    def refresh_master_tables() -> Dict[str, int]:
        data = fetch_live_ezaccount_master_data()
        if not any(data.get(key) for key in ("vendors", "gl_accounts", "cost_centers")):
            raise RuntimeError("EZAccount 선택지를 가져오지 못했습니다. 로그인 및 PR 생성 화면을 확인하세요.")
        conn = get_db_connection()
        conn.execute("DELETE FROM vendors")
        vendor_count = sync_master_source(conn, "vendor", data.get("vendors", []))
        gl_count = sync_choice_source(conn, "gl_accounts", data.get("gl_accounts", []))
        cost_center_count = sync_choice_source(conn, "cost_centers", data.get("cost_centers", []))
        conn.execute(
            "INSERT INTO master_refresh_state(id, refreshed_at) VALUES (1, ?) ON CONFLICT(id) DO UPDATE SET refreshed_at = excluded.refreshed_at",
            (utc_now(),),
        )
        conn.commit()
        conn.close()
        return {"vendors": vendor_count, "gl_accounts": gl_count, "cost_centers": cost_center_count}

    @app.route("/api/masters")
    def masters_api():
        conn = get_db_connection()
        vendors = conn.execute("SELECT vendor_code AS code, vendor_name AS name FROM vendors ORDER BY vendor_name").fetchall()
        gl_accounts = conn.execute("SELECT gl_code AS code, gl_name AS name FROM gl_accounts ORDER BY gl_code").fetchall()
        cost_centers = conn.execute("SELECT cost_center_code AS code, cost_center_name AS name FROM cost_centers ORDER BY cost_center_code").fetchall()
        refresh_state = conn.execute("SELECT refreshed_at FROM master_refresh_state WHERE id = 1").fetchone()
        conn.close()
        if refresh_state is None:
            return jsonify({"vendors": [], "gl_accounts": [], "cost_centers": [], "refreshed_at": None})
        return jsonify({
            "vendors": [dict(row) for row in vendors],
            "gl_accounts": [dict(row) for row in gl_accounts],
            "cost_centers": [dict(row) for row in cost_centers],
            "refreshed_at": refresh_state["refreshed_at"],
        })

    @app.route("/api/masters/refresh", methods=["POST"])
    def refresh_masters_api():
        try:
            stats = refresh_master_tables()
            return jsonify({"ok": True, "updated": stats})
        except Exception as exc:
            logging.exception("EZAccount master refresh failed")
            return jsonify({"ok": False, "error": str(exc)}), 502

    @app.route("/api/masters/import", methods=["POST"])
    def import_master_data():
        file = request.files.get("file")
        kind = request.form.get("kind", "vendor")
        if file is None or file.filename == "":
            return jsonify({"ok": False, "error": "업로드할 CSV 또는 Excel 파일을 선택하세요."}), 400

        try:
            rows = read_master_file(file)
            if kind == "vendor":
                parsed = normalize_master_rows(rows, "vendor")
                conn = get_db_connection()
                for item in parsed:
                    conn.execute(
                        "INSERT OR IGNORE INTO vendors(vendor_code, vendor_name) VALUES (?, ?)",
                        (item["code"], item["name"]),
                    )
                conn.commit()
                conn.close()
                return jsonify({"ok": True, "kind": "vendor", "count": len(parsed)})

            if kind == "material":
                parsed = normalize_master_rows(rows, "material")
                conn = get_db_connection()
                for item in parsed:
                    conn.execute(
                        "INSERT OR IGNORE INTO materials(material_code, material_name) VALUES (?, ?)",
                        (item["code"], item["name"]),
                    )
                conn.commit()
                conn.close()
                return jsonify({"ok": True, "kind": "material", "count": len(parsed)})

            return jsonify({"ok": False, "error": "kind는 vendor 또는 material 이어야 합니다."}), 400
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400

    @app.route("/api/requests")
    def requests_api():
        start_rejection_monitor()
        conn = get_db_connection()
        rows = conn.execute(
            "SELECT * FROM purchase_requests ORDER BY id DESC"
        ).fetchall()
        conn.close()
        return jsonify({"items": [dict(r) for r in rows]})

    @app.route("/api/pr", methods=["POST"])
    def create_pr():
        payload = request.get_json(silent=True) or {}
        try:
            form_items = payload.get("items") or json.loads(request.form.get("items") or "[]")
            attachment_manifest = json.loads(request.form.get("attachment_manifest") or "[]")
        except json.JSONDecodeError:
            return jsonify({"ok": False, "error": "Invalid item or attachment metadata."}), 400
        if not form_items:
            return jsonify({"ok": False, "error": "No items provided"}), 400

        request_id = None
        conn = get_db_connection()
        created_at = utc_now()
        total_amount = sum(float(item_total(item)) for item in form_items)
        cur = conn.execute(
            "INSERT INTO purchase_requests (status, revision, total_amount, created_at, updated_at, notes) VALUES (?, 0, ?, ?, ?, ?)",
            ("draft", total_amount, created_at, created_at, "Created from web form"),
        )
        request_id = cur.lastrowid

        for idx, item in enumerate(form_items, start=1):
            qty = float(item.get("quantity") or 0)
            unit_price = float(item.get("unit_price") or item.get("unitPrice") or 0)
            total_price = float(item.get("total_price") or item.get("totalPrice") or (qty * unit_price))
            material_code = str(item.get("material_code") or item.get("materialCode") or "").strip()
            material_name = str(item.get("material_name") or item.get("materialName") or "").strip()
            if not material_code and not material_name:
                material_code = ""
                material_name = ""
            conn.execute(
                "INSERT INTO pr_items (request_id, quote_group, item_name, purchase_reason, quantity, unit_price, total_price, item_link, vendor_code, vendor_name, material_code, material_name, material_group, manufacturing_part_no, gl_account, cost_center, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    request_id,
                    str(item.get("quote_group") or item.get("quoteGroup") or f"quote-{idx}"),
                    str(item.get("item_name") or item.get("itemName") or ""),
                    str(item.get("purchase_reason") or item.get("purchaseReason") or ""),
                    qty,
                    unit_price,
                    total_price,
                    str(item.get("item_link") or item.get("itemLink") or ""),
                    str(item.get("vendor_code") or ""),
                    str(item.get("vendor_name") or ""),
                    material_code,
                    material_name,
                    str(item.get("material_group") or item.get("materialGroup") or ""),
                    str(item.get("manufacturing_part_no") or item.get("manufacturingPartNo") or ""),
                    str(item.get("gl_account") or item.get("glAccount") or ""),
                    str(item.get("cost_center") or item.get("costCenter") or ""),
                    created_at,
                ),
            )
        conn.commit()
        conn.close()
        save_uploaded_files(request_id, request.files.getlist("files"), attachment_manifest)

        batches = suggest_batches(form_items)
        draft = build_ezaccount_draft(request_id)
        create_revision(request_id, 0, "Initial draft created")
        return jsonify({"ok": True, "request_id": request_id, "batches": batches, "draft": draft})

    @app.route("/api/pr/<int:request_id>/draft")
    def draft_api(request_id: int):
        try:
            return jsonify(build_ezaccount_draft(request_id))
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 404

    @app.route("/api/pr/<int:request_id>/ezaccount", methods=["POST"])
    def ezaccount_submit_api(request_id: int):
        try:
            result = run_ezaccount_automation(request_id)
            if result.get("ok"):
                document_no = str(result.get("document_no") or "").strip()
                conn = get_db_connection()
                conn.execute(
                    "UPDATE purchase_requests SET status = ?, ez_doc_no = ?, ez_created_at = ?, ez_status = ?, updated_at = ? WHERE id = ?",
                    (
                        "saved",
                        document_no or None,
                        utc_now(),
                        "Saved - awaiting manual review and submission",
                        utc_now(),
                        request_id,
                    ),
                )
                conn.commit()
                conn.close()
            return jsonify(result)
        except (RuntimeError, ValueError) as exc:
            return jsonify({"ok": False, "error": str(exc), "request_id": request_id}), 400
        except Exception as exc:  # pragma: no cover
            return jsonify({"ok": False, "error": f"EZAccount automation failed: {exc}", "request_id": request_id}), 500

    @app.route("/api/pr/<int:request_id>/revise", methods=["POST"])
    def revise_pr(request_id: int):
        payload = request.get_json(silent=True) or {}
        revision = int(payload.get("revision", 0)) + 1
        summary = payload.get("summary") or f"Revision {revision} update"
        create_revision(request_id, revision, summary)
        return jsonify({"ok": True, "revision": revision, "summary": summary})

    return app


app = create_app()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
