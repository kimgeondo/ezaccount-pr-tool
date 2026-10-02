from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEADER_COLUMNS = [
    "USE",
    "기안자",
    "작성날짜",
    "자산명 (Name of Asset)",
    "문서번호 (Document Number)",
    "업체명 (Vendor)",
    "사양 및 모델 번호 (Specification / Model Number)",
    "수량 (Quantity)",
    "P.O NUMBER",
    "입고유무",
    "용도",
    "비고",
]

TITLE_FILL = PatternFill(fill_type="solid", fgColor="FFF9C4")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="C5E1A5")
BORDER_STYLE = Border(
    left=Side(border_style="thin", color="000000"),
    right=Side(border_style="thin", color="000000"),
    top=Side(border_style="thin", color="000000"),
    bottom=Side(border_style="thin", color="000000"),
)


def create_template(path, title_text="2023-2024 보전 구매 내역", department_text="부서, AJINUSA MAINTENANCE"):
    template_path = Path(path)
    template_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PR 누적"

    total_columns = len(HEADER_COLUMNS)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columns)
    header_cell = sheet.cell(row=1, column=1)
    header_cell.value = title_text
    header_cell.font = Font(bold=True, size=14)
    header_cell.alignment = Alignment(horizontal="center", vertical="center")
    header_cell.fill = TITLE_FILL
    header_cell.border = BORDER_STYLE

    sheet.cell(row=2, column=1, value=department_text)
    sheet.cell(row=2, column=1).alignment = Alignment(horizontal="left", vertical="center")

    sheet.row_dimensions[1].height = 26
    sheet.row_dimensions[2].height = 18

    for index, column_name in enumerate(HEADER_COLUMNS, start=1):
        cell = sheet.cell(row=4, column=index, value=column_name)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_STYLE

    sheet.auto_filter.ref = f"A4:{chr(ord('A') + total_columns - 1)}4"
    sheet.freeze_panes = "A5"

    widths = [14, 12, 15, 24, 18, 20, 30, 10, 16, 12, 24, 32]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=4, column=index).column_letter].width = width

    for row in range(1, 5):
        for col in range(1, total_columns + 1):
            sheet.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="center")
            sheet.cell(row=row, column=col).border = BORDER_STYLE

    workbook.save(template_path)
    return str(template_path)
