from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side

from .excel_template import HEADER_COLUMNS, create_template


def _ensure_workbook(output_file, template_file):
    output_path = Path(output_file)
    if not output_path.exists():
        if not Path(template_file).exists():
            create_template(template_file)
        workbook = load_workbook(template_file)
        workbook.save(output_path)

    return load_workbook(output_path)


def append_rows(output_file, template_file, rows):
    workbook = _ensure_workbook(output_file, template_file)
    sheet = workbook.active
    start_row = max(sheet.max_row + 1, 5)
    border = Border(
        left=Side(border_style="thin", color="000000"),
        right=Side(border_style="thin", color="000000"),
        top=Side(border_style="thin", color="000000"),
        bottom=Side(border_style="thin", color="000000"),
    )
    alignment = Alignment(wrap_text=True, vertical="top")

    for index, row in enumerate(rows, start=0):
        row_number = start_row + index
        for col_idx, column_name in enumerate(HEADER_COLUMNS, start=1):
            value = row.get(column_name, "")
            cell = sheet.cell(row=row_number, column=col_idx, value=value)
            cell.alignment = alignment
            cell.border = border

    workbook.save(output_file)
    return len(rows)
