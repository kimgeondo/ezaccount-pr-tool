# src/main.py
from __future__ import annotations

import json
import logging
import os
import re
import shutil
import tempfile
import time
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from typing import Dict, List, Tuple
from openpyxl.styles import Border, Side

from .client import ODataClient

# =========================
# Excel Columns (visible)
# =========================
COLUMNS: List[str] = [
    "사용처 (USE)",
    "기안자",
    "작성날짜",
    "자산명 (Name of Asset)",  # ZTXZ01
    "문서번호 (Document Number)",
    "업체명 (Vendor)",          # NAME1
    "사양 및 모델 번호 (Specification / Model Number)",  # MFRPN or MATNR
    "수량 (Quantity)",          # MENGE
    "Unit Price",               # NETPR
    "Total Price",              # NETWR
    "P.O NUMBER",               # EBELN
    "입고유무",                 # blank
    "용도",                     # ZTEXT3 numbering -> line index mapping
    "비고"                      # ALWAYS BLANK (USD removed)
]

DOC_COL = 5                 # Document Number column (E)
PO_COL = 11                 # PO NUMBER column (K)
HIDDEN_KEY_NAME = "_LINE_KEY"  # hidden key column name (added at end)


# =========================
# Logging
# =========================
def setup_logging(log_dir: Path) -> None:
    log_dir.mkdir(parents=True, exist_ok=True)
    log_file = log_dir / "latest.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        handlers=[logging.FileHandler(log_file, encoding="utf-8"), logging.StreamHandler()],
    )


# =========================
# State
# =========================
def load_state(state_file: Path) -> Dict[str, Any]:
    if not state_file.exists():
        return {"last_sync": None, "seen_keys": [], "rejected_seen_keys": []}
    state = json.loads(state_file.read_text(encoding="utf-8"))
    state.setdefault("seen_keys", [])
    state.setdefault("rejected_seen_keys", [])
    return state


def save_state(state_file: Path, state: Dict[str, Any]) -> None:
    state_file.parent.mkdir(parents=True, exist_ok=True)
    state_file.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")


def now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def parse_iso(s: str) -> datetime:
    return datetime.fromisoformat(s.replace("Z", "+00:00"))


# =========================
# Helpers
# =========================
def odata_results(nav: Any) -> List[Dict[str, Any]]:
    if isinstance(nav, dict) and isinstance(nav.get("results"), list):
        return nav["results"]
    return []


def normalize_yyyymmdd(s: Any) -> str:
    txt = str(s or "").strip()
    if re.fullmatch(r"\d{8}", txt):
        return f"{txt[0:4]}-{txt[4:6]}-{txt[6:8]}"
    return txt


def year_from_reqdt(reqdt: Any) -> str:
    s = str(reqdt or "").strip()
    if len(s) >= 4 and s[:4].isdigit():
        return s[:4]
    m = re.match(r"^(\d{4})-\d{2}-\d{2}$", s)
    return m.group(1) if m else "Unknown"


def to_number(x: Any) -> Any:
    if x is None:
        return ""
    if isinstance(x, (int, float)):
        return x
    txt = str(x).strip().replace(",", "")
    if txt == "":
        return ""
    try:
        return float(txt)
    except Exception:
        return x


def safe_save_workbook(wb: Workbook, output_path: Path, retries: int = 8, delay: float = 0.8) -> None:
    """
    Save with retry (Excel file might be open/locked).
    """
    output_path = Path(output_path)

    for _ in range(retries):
        try:
            wb.save(output_path)
            return
        except PermissionError:
            time.sleep(delay)

    fd, tmp_name = tempfile.mkstemp(
        prefix=f"{output_path.stem}.",
        suffix=output_path.suffix,
        dir=output_path.parent,
    )
    os.close(fd)
    tmp_path = Path(tmp_name)

    wb.save(tmp_path)
    try:
        os.replace(tmp_path, output_path)
    except PermissionError as e:
        raise PermissionError(
            f"파일이 잠겨서 저장 불가: {output_path}\n"
            f"엑셀 닫고 재실행. 임시파일: {tmp_path}"
        ) from e


def publish_shared_workbook(
    source_path: Path,
    shared_dir: Path,
    shared_filename: str,
    retries: int = 8,
    delay: float = 0.8,
) -> Path:
    source_path = Path(source_path)
    shared_dir = Path(shared_dir).expanduser()
    if not source_path.is_file():
        raise FileNotFoundError(f"완료 승인 엑셀 원본을 찾을 수 없습니다: {source_path}")

    shared_dir.mkdir(parents=True, exist_ok=True)
    target_path = shared_dir / shared_filename
    if source_path.resolve() == target_path.resolve():
        return target_path

    fd, temp_name = tempfile.mkstemp(
        prefix=f"{target_path.stem}.",
        suffix=".tmp.xlsx",
        dir=shared_dir,
    )
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        shutil.copy2(source_path, temp_path)
        for attempt in range(retries):
            try:
                shutil.copyfile(temp_path, target_path)
                return target_path
            except PermissionError:
                if attempt + 1 == retries:
                    raise
                time.sleep(delay)
    finally:
        if temp_path.exists():
            temp_path.unlink()

    return target_path


# =========================
# Filters
# =========================
def is_accounting_completed(pr: Dict[str, Any]) -> bool:
    # confirmed by your data: BSTAT_NAME == "Accounting Completed"
    return str(pr.get("BSTAT_NAME", "")).strip().lower() == "accounting completed"


def is_rejected(pr: Dict[str, Any]) -> bool:
    """PR이 반려 상태인지 판별한다."""
    status = str(pr.get("BSTAT_NAME", pr.get("CSTAT_NAME", ""))).strip()
    normalized = re.sub(r"\s+", " ", status).lower()
    return "reject" in normalized


# =========================
# Item Description parsing (ZTEXT3)
# =========================
def parse_item_description(text: str) -> Dict[int, str]:
    """
    1) ...
    3, 4) ...
    3-5) ... / 3~5) ...
    -> {1:'..',3:'..',4:'..',5:'..'}
    """
    if not text:
        return {}

    mapping: Dict[int, str] = {}
    lines = [ln.strip() for ln in str(text).splitlines() if ln.strip()]

    for ln in lines:
        m = re.match(r"^(\d+(?:\s*[,~-]\s*\d+)*)\)\s*(.+)$", ln)
        if not m:
            continue

        nums_part = m.group(1).strip()
        desc = m.group(2).strip()

        nums: List[int] = []
        parts = [p.strip() for p in re.split(r"\s*,\s*", nums_part)]
        for p in parts:
            rng = re.split(r"\s*[-~]\s*", p)
            if len(rng) == 2 and rng[0].isdigit() and rng[1].isdigit():
                a, b = int(rng[0]), int(rng[1])
                step = 1 if a <= b else -1
                nums.extend(list(range(a, b + step, step)))
            elif p.isdigit():
                nums.append(int(p))

        for n in nums:
            mapping[n] = desc

    return mapping


# =========================
# Approval: REQ approver (fallback)
# =========================
def extract_req_approver(appr_lines: List[Dict[str, Any]]) -> str:
    if not appr_lines:
        return ""

    step_keys = ["APPR_STEP", "STEP", "APP_STEP", "APPRSTEP", "ApprovalStep", "ZSTEP"]
    approver_keys = ["APPROVER", "APPROVER_NAME", "Approver", "APPR_NAME", "NAME", "TEXT", "ENAME"]

    def pick(d: Dict[str, Any], keys: List[str]) -> str:
        for k in keys:
            v = d.get(k)
            if v not in (None, ""):
                return str(v).strip()
        return ""

    for row in appr_lines:
        step = pick(row, step_keys)
        if step == "REQ":
            return pick(row, approver_keys)

    return ""


# =========================
# Output file naming (rename-in-place)
# =========================
def todays_filename(fmt: str) -> str:
    MM = datetime.now().strftime("%m")
    DD = datetime.now().strftime("%d")
    YYYY = datetime.now().strftime("%Y")
    return fmt.format(MM=MM, DD=DD, YYYY=YYYY)


from typing import Optional  # 파일 상단 import들에 Optional이 없으면 추가

def pick_existing_file(output_dir: Path, pattern: str) -> Optional[Path]:
    """
    output_dir에서 pattern에 매칭되는 파일 중 가장 최근 수정된 파일 1개 반환
    임시 .tmp.xlsx 파일은 제외합니다.
    없으면 None
    """
    files = [p for p in output_dir.glob(pattern) if not p.name.endswith(".tmp.xlsx")]
    if not files:
        return None
    files.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return files[0]


# =========================
# Workbook / Sheet
# =========================
def ensure_workbook(path: Path) -> Workbook:
    if path.exists():
        return load_workbook(path)
    wb = Workbook()
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1:
        wb.remove(wb["Sheet"])
    return wb


def ensure_year_sheet(
    wb: Workbook,
    sheet_name: str,
    title_text: str,
    header_row: int,
    department_name: str
) -> Worksheet:
    """
    Create year sheet if not exists, and ensure hidden key column.
    """
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(sheet_name)

        title_fill = PatternFill("solid", fgColor="FFF2CC")
        header_fill = PatternFill("solid", fgColor="C6EFCE")
        bold = Font(bold=True)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
        t = ws.cell(row=1, column=1, value=title_text)
        t.fill = title_fill
        t.font = Font(bold=True, size=14)
        t.alignment = center
        ws.row_dimensions[1].height = 24

        ws.cell(row=2, column=1, value="부서").font = bold
        ws.cell(row=2, column=2, value=department_name)

        thin = Side(style="thin", color="D9D9D9")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for i, col_name in enumerate(COLUMNS, start=1):
            c = ws.cell(row=header_row, column=i, value=col_name)
            c.fill = header_fill
            c.font = bold
            c.alignment = center
            c.border = border

            letter = get_column_letter(i)
            if col_name in ("사양 및 모델 번호 (Specification / Model Number)", "용도", "비고"):
                ws.column_dimensions[letter].width = 30
            elif col_name in ("업체명 (Vendor)", "자산명 (Name of Asset)"):
                ws.column_dimensions[letter].width = 22
            else:
                ws.column_dimensions[letter].width = 14

        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(COLUMNS))}{header_row}"
        ws.freeze_panes = f"A{header_row + 1}"

    # hidden key column at end
    key_col = len(COLUMNS) + 1
    if ws.cell(row=header_row, column=key_col).value != HIDDEN_KEY_NAME:
        ws.cell(row=header_row, column=key_col, value=HIDDEN_KEY_NAME)
    ws.column_dimensions[get_column_letter(key_col)].hidden = True
    ws.column_dimensions[get_column_letter(key_col)].width = 2

    return ws


def find_next_row(ws: Worksheet, header_row: int) -> int:
    nxt = ws.max_row + 1
    min_row = header_row + 1
    return nxt if nxt >= min_row else min_row


def build_key_index(ws: Worksheet, header_row: int) -> Dict[str, int]:
    key_col = len(COLUMNS) + 1
    idx: Dict[str, int] = {}
    start = header_row + 1
    for r in range(start, ws.max_row + 1):
        k = ws.cell(r, key_col).value
        if k is None:
            continue
        ks = str(k).strip()
        if ks:
            idx[ks] = r
    return idx


def append_or_update_rows(
    ws: Worksheet,
    header_row: int,
    rows_with_key: List[Tuple[List[Any], str]],
    key_to_row: Dict[str, int],
) -> Tuple[int, int]:
    """
    - If key exists: update PO only if empty and new value exists.
    - If key not exists: append row and store hidden key.
    Returns: (appended_count, updated_po_count)
    """
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(vertical="top", wrap_text=True)

    key_col = len(COLUMNS) + 1
    appended = 0
    updated = 0
    next_row = find_next_row(ws, header_row)

    for row_vals, line_key in rows_with_key:
        if line_key in key_to_row:
            r = key_to_row[line_key]
            existing_po = ws.cell(r, PO_COL).value
            new_po = row_vals[PO_COL - 1]
            if (existing_po is None or str(existing_po).strip() == "") and (new_po is not None and str(new_po).strip() != ""):
                ws.cell(r, PO_COL, value=new_po)
                updated += 1
            continue

        # append new
        for c, val in enumerate(row_vals, start=1):
            cell = ws.cell(next_row, c, value=val)
            cell.border = border
            cell.alignment = wrap

        ws.cell(next_row, key_col, value=line_key)
        key_to_row[line_key] = next_row

        next_row += 1
        appended += 1

    return appended, updated


def post_format_grouping(ws: Worksheet, header_row: int) -> None:
    """
    - 문서번호(E열) 동일 문서 구간 병합(문서번호 컬럼만)
    - 각 문서 구간 '시작행 위'에 굵은선(Top), '마지막행 아래'에 굵은선(Bottom)
    - 병합 셀로 인해 문서번호 값이 빈칸이 되는 문제를 carry-forward로 해결
    """
    first = header_row + 1
    last = ws.max_row
    if last < first:
        return

    # 1) 기존 문서번호 병합 해제 (문서번호 컬럼만)
    for rng in list(ws.merged_cells.ranges):
        if rng.min_col == DOC_COL and rng.max_col == DOC_COL and rng.min_row >= first:
            ws.unmerge_cells(str(rng))

    # 2) carry-forward로 doc 값 채우기 (빈칸이면 바로 위 doc 유지)
    doc_by_row: Dict[int, str] = {}
    current_doc = ""
    for r in range(first, last + 1):
        v = ws.cell(r, DOC_COL).value
        s = str(v).strip() if v is not None else ""
        if s:
            current_doc = s
        doc_by_row[r] = current_doc

    # 3) 기존 굵은선 초기화 (top/bottom을 thin으로 되돌림)
    thin = Side(style="thin", color="D9D9D9")
    for r in range(first, last + 1):
        for c in range(1, len(COLUMNS) + 1):
            cell = ws.cell(r, c)
            b = cell.border
            cell.border = Border(
                left=b.left, right=b.right,
                top=thin, bottom=thin,          # ⭐ top/bottom 초기화
                diagonal=b.diagonal, diagonal_direction=b.diagonal_direction,
                outline=b.outline, vertical=b.vertical, horizontal=b.horizontal
            )

    # 4) 연속 동일 doc 그룹 만들기
    groups: List[Tuple[str, int, int]] = []
    r = first
    while r <= last:
        doc = doc_by_row.get(r, "")
        if not doc:
            r += 1
            continue
        start = r
        while r + 1 <= last and doc_by_row.get(r + 1, "") == doc:
            r += 1
        end = r
        groups.append((doc, start, end))
        r += 1

    # 5) 문서번호 컬럼 병합
    for doc, start, end in groups:
        if end > start:
            ws.merge_cells(start_row=start, start_column=DOC_COL, end_row=end, end_column=DOC_COL)

    # 6) 문서별 시작행 TOP 굵은선 + 마지막행 BOTTOM 굵은선
    thick = Side(style="thick", color="000000")

    for doc, start, end in groups:
        # ⭐ start행: Top thick
        for c in range(1, len(COLUMNS) + 1):
            cell = ws.cell(start, c)
            b = cell.border
            cell.border = Border(
                left=b.left, right=b.right,
                top=thick, bottom=b.bottom,     # TOP만 굵게
                diagonal=b.diagonal, diagonal_direction=b.diagonal_direction,
                outline=b.outline, vertical=b.vertical, horizontal=b.horizontal
            )

        # ⭐ end행: Bottom thick
        for c in range(1, len(COLUMNS) + 1):
            cell = ws.cell(end, c)
            b = cell.border
            cell.border = Border(
                left=b.left, right=b.right,
                top=b.top, bottom=thick,        # BOTTOM만 굵게
                diagonal=b.diagonal, diagonal_direction=b.diagonal_direction,
                outline=b.outline, vertical=b.vertical, horizontal=b.horizontal
            )


# =========================
# Build rows per PR
# =========================
def make_line_key(bukrs: str, bnumb: str, gjahr: str, zbnfpo: Any) -> str:
    return f"{bukrs}|{bnumb}|{gjahr}|{str(zbnfpo).strip()}"


def build_rows_for_document(pr_row: Dict[str, Any], header: Dict[str, Any]) -> Tuple[List[List[Any]], List[Dict[str, Any]]]:
    """
    Uses:
      Header purpose: ZTEXT3 (numbered)
      Items:
        ZTXZ01, NAME1, MENGE, NETPR, NETWR, MFRPN, MATNR, EBELN, ZBNFPO
    """
    use_dept = str(pr_row.get("REQDP_NAME", "")).strip()
    doc_no = str(pr_row.get("BNUMB", "")).strip()
    req_date = normalize_yyyymmdd(pr_row.get("REQDT", ""))

    appr_lines = odata_results(header.get("ApprLineNavi"))
    requester = extract_req_approver(appr_lines) or str(pr_row.get("REQID_NAME", "")).strip()

    purpose_text = str(header.get("ZTEXT3", "") or "")
    purpose_map = parse_item_description(purpose_text)

    items = odata_results(header.get("PRItemNavi"))

    def zbnfpo_int(it: Dict[str, Any]) -> int:
        try:
            return int(str(it.get("ZBNFPO", "")).strip() or "0")
        except Exception:
            return 0

    items_sorted = sorted(items, key=zbnfpo_int)

    rows: List[List[Any]] = []
    for idx, it in enumerate(items_sorted, start=1):
        asset_name = str(it.get("ZTXZ01", "")).strip()
        vendor = str(it.get("NAME1", "")).strip()
        qty = to_number(it.get("MENGE", ""))
        unit_price = to_number(it.get("NETPR", ""))
        total_price = to_number(it.get("NETWR", ""))
        spec = str(it.get("MFRPN", "")).strip() or str(it.get("MATNR", "")).strip()
        po_no = str(it.get("EBELN", "")).strip()
        purpose = purpose_map.get(idx, "")

        # USD 제거: always blank
        remark = ""

        row = [
            use_dept,
            requester,
            req_date,
            asset_name,
            doc_no,
            vendor,
            spec,
            qty,
            unit_price,
            total_price,
            po_no,
            "",
            purpose,
            remark
        ]
        rows.append(row)

    return rows, items_sorted


def write_yearly_prs_to_output(
    year_to_prs: Dict[str, List[Dict[str, Any]]],
    output_dir: Path,
    filename_format: str,
    filename_glob: str,
    title_fmt: str,
    department_name: str,
    header_row: int,
    client: ODataClient,
    seen_keys: Set[str],
) -> Tuple[int, int, Path]:
    """주어진 상태군(year_to_prs)을 기준으로 엑셀 파일을 생성/업데이트한다."""
    desired_name = todays_filename(filename_format)
    desired_path = output_dir / desired_name
    base_path = desired_path

    files = [p for p in output_dir.glob(filename_glob) if not p.name.endswith('.tmp.xlsx')]
    files.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    sheet_to_file: Dict[str, Path] = {}
    for p in files:
        try:
            wb_tmp = load_workbook(p, read_only=True)
            for s in wb_tmp.sheetnames:
                if s not in sheet_to_file:
                    sheet_to_file[s] = p
            wb_tmp.close()
        except Exception:
            continue

    appended_total = 0
    updated_po_total = 0
    key_index_cache: Dict[Tuple[str, str], Dict[str, int]] = {}

    for year, prs in sorted(year_to_prs.items()):
        sheet_name = str(year)
        title_text = title_fmt.format(year=year)
        target_path = sheet_to_file.get(sheet_name, base_path)

        wb = ensure_workbook(target_path)
        ws = ensure_year_sheet(wb, sheet_name, title_text, header_row, department_name)

        cache_key = (str(target_path.resolve()), sheet_name)
        if cache_key in key_index_cache:
            key_to_row = key_index_cache[cache_key]
        else:
            key_to_row = build_key_index(ws, header_row)
            key_index_cache[cache_key] = key_to_row

        for pr in prs:
            bukrs = str(pr.get("BUKRS", "")).strip()
            bnumb = str(pr.get("BNUMB", "")).strip()
            gjahr = str(pr.get("GJAHR", "")).strip()
            if not (bukrs and bnumb and gjahr):
                continue

            header = client.fetch_header(bukrs, bnumb, gjahr)
            if not header:
                continue

            rows, items_sorted = build_rows_for_document(pr, header)
            if not rows:
                continue

            rows_with_key: List[Tuple[List[Any], str]] = []
            for i, row_vals in enumerate(rows):
                zbnfpo = items_sorted[i].get("ZBNFPO", "") if i < len(items_sorted) else (i + 1)
                line_key = make_line_key(bukrs, bnumb, gjahr, zbnfpo)
                rows_with_key.append((row_vals, line_key))

            rows_with_key.sort(key=lambda x: str(x[0][DOC_COL - 1]))

            appended, updated = append_or_update_rows(ws, header_row, rows_with_key, key_to_row)
            appended_total += appended
            updated_po_total += updated

            for _, k in rows_with_key:
                seen_keys.add(k)

        post_format_grouping(ws, header_row)
        safe_save_workbook(wb, target_path)
        wb.close()

    try:
        if base_path.resolve() != desired_path.resolve() and base_path.exists():
            if desired_path.exists():
                try:
                    desired_path.unlink()
                except PermissionError:
                    raise PermissionError(
                        f"오늘 파일({desired_path})이 열려 있어 이름 변경 불가. 엑셀 닫고 재실행."
                    )
            try:
                os.replace(base_path, desired_path)
            except PermissionError as e:
                raise PermissionError(
                    f"파일명 변경 실패(파일이 열려있을 수 있음):\n"
                    f"FROM: {base_path}\nTO:   {desired_path}\n"
                    f"엑셀 닫고 다시 실행하세요."
                ) from e
    except Exception:
        raise

    return appended_total, updated_po_total, desired_path


# =========================
# Main
# =========================
def main() -> None:
    load_dotenv()

    project_root = Path(__file__).resolve().parents[1]
    config_path = project_root / "config" / "config.json"
    cfg = json.loads(config_path.read_text(encoding="utf-8"))

    setup_logging(project_root / "logs")
    logging.info("Starting ezaccount PR sync")

    state_file = project_root / cfg.get("state_file", "state/state.json")

    output_dir = project_root / str(cfg.get("output_dir", "output"))
    output_dir.mkdir(parents=True, exist_ok=True)

    fmt = str(cfg.get("output_filename_format", "MT_PRS_PR_LIST_{MM}_{DD}_{YYYY}.xlsx"))
    pattern = str(cfg.get("output_filename_glob", "MT_PRS_PR_LIST_*.xlsx"))
    rejected_fmt = str(cfg.get("rejected_output_filename_format", "MT_PRS_PR_REJECTED_{MM}_{DD}_{YYYY}.xlsx"))
    rejected_pattern = str(cfg.get("rejected_output_filename_glob", "MT_PRS_PR_REJECTED_*.xlsx"))

    header_row = int(cfg.get("data_start_row", 5))
    department_name = str(cfg.get("department_name", "AJINUSA MAINTENANCE"))
    title_fmt = str(cfg.get("sheet_title_format", "{year} 보전 구매 내역"))
    lookback_days = int(cfg.get("po_update_lookback_days", 90))

    state = load_state(state_file)
    last_sync: Optional[str] = state.get("last_sync")
    seen: Set[str] = set(state.get("seen_keys", []))
    rejected_seen: Set[str] = set(state.get("rejected_seen_keys", []))

    # lookback for PO updates
    fetch_since: Optional[str] = None
    if last_sync:
        try:
            dt = parse_iso(last_sync) - timedelta(days=lookback_days)
            fetch_since = dt.replace(microsecond=0).isoformat().replace("+00:00", "Z")
        except Exception:
            fetch_since = last_sync

    client = ODataClient(str(config_path))

    # Support multiple requester IDs (config.requester_ids)
    pr_list: List[Dict[str, Any]] = []
    requester_ids = cfg.get("requester_ids")
    reqdp_filter = cfg.get("reqdp_filter")

    if requester_ids and isinstance(requester_ids, list) and len(requester_ids) > 0:
        for reqid in requester_ids:
            batch = client.fetch_list(fetch_since, reqid_override=reqid, reqdp_override=reqdp_filter)
            pr_list.extend(batch)
    else:
        pr_list = client.fetch_list(fetch_since)

    logging.info("Fetched %d records from PRListSet (lookback %d days)", len(pr_list), lookback_days)

    # Group PRs by year to ensure each year's sheet is updated in the correct workbook
    year_to_prs: Dict[str, List[Dict[str, Any]]] = {}
    rejected_year_to_prs: Dict[str, List[Dict[str, Any]]] = {}
    for pr in pr_list:
        if is_accounting_completed(pr):
            year = year_from_reqdt(pr.get("REQDT"))
            year_to_prs.setdefault(year, []).append(pr)
        elif is_rejected(pr):
            year = year_from_reqdt(pr.get("REQDT"))
            rejected_year_to_prs.setdefault(year, []).append(pr)

    appended_total = 0
    updated_po_total = 0

    completed_path = output_dir / todays_filename(fmt)
    if year_to_prs:
        appended, updated, completed_path = write_yearly_prs_to_output(
            year_to_prs,
            output_dir,
            fmt,
            pattern,
            title_fmt,
            department_name,
            header_row,
            client,
            seen,
        )
        appended_total += appended
        updated_po_total += updated

    rejected_appended_total = 0
    rejected_updated_total = 0
    rejected_path = output_dir / todays_filename(rejected_fmt)
    if rejected_year_to_prs:
        rejected_appended_total, rejected_updated_total, rejected_path = write_yearly_prs_to_output(
            rejected_year_to_prs,
            output_dir,
            rejected_fmt,
            rejected_pattern,
            title_fmt,
            department_name,
            header_row,
            client,
            rejected_seen,
        )

    shared_dir_value = str(cfg.get("shared_output_dir", "")).strip()
    if shared_dir_value:
        latest_workbook = pick_existing_file(output_dir, pattern)
        if latest_workbook is None:
            logging.warning("No completed workbook found to publish to the shared folder.")
        else:
            shared_filename = str(cfg.get("shared_output_filename", "MT_PRS_PR_LIST_SHARED.xlsx"))
            shared_path = publish_shared_workbook(latest_workbook, Path(shared_dir_value), shared_filename)
            logging.info("Shared workbook updated: %s (source: %s)", shared_path, latest_workbook)

    # update state
    state["last_sync"] = now_iso()
    state["seen_keys"] = list(seen)
    state["rejected_seen_keys"] = list(rejected_seen)
    save_state(state_file, state)

    logging.info("Appended new rows: %d", appended_total)
    logging.info("Updated PO cells: %d", updated_po_total)
    logging.info("Final output: %s", str(completed_path if year_to_prs else output_dir / todays_filename(fmt)))
    if rejected_year_to_prs:
        logging.info("Rejected appended new rows: %d", rejected_appended_total)
        logging.info("Rejected updated PO cells: %d", rejected_updated_total)
        logging.info("Rejected output: %s", str(rejected_path))
    logging.info("Sync completed. last_sync updated.")


if __name__ == "__main__":
    main()